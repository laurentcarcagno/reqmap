"""Excel export — fills the HMA-C-049 requirements compliance matrix template."""

import re
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from app.config import get_settings
from app.models.document import Document
from app.models.node import Node
from app.schemas.export import DocumentRef, ExportSettings
from app.services.markdown_cells import Picture, load_picture, render_table, split_segments, table_as_text

TEMPLATE_NAME = "HMA-C-049_EdC_Matrice-Conformite-Exigences.xlsx"

# REQUIREMENTS sheet layout: header on row 7, template example rows 8-11
_FIRST_ROW = 8
_EXAMPLE_LAST_ROW = 11
_TITLE_STYLE_ROW = 8  # bold title row
_REQUIREMENT_STYLE_ROW = 10  # requirement row
_COLUMNS = range(1, 7)  # A..F: Title, Identifier, Description, Applicability, Conformity, Comment
_MAX_CELL_CHARS = 32_000  # Excel limit is 32,767

# DA-DR sheet: first data row and capacity of each table in the template
_AD_FIRST, _AD_CAPACITY = 10, 4
_RD_FIRST, _RD_CAPACITY = 19, 4
_ACRONYM_FIRST, _ACRONYM_CAPACITY = 28, 8

# Row height estimate for a description cell that carries pictures (Arial 10, column C)
_LINE_PX = 17
_CHARS_PER_LINE = 80
_GAP_PX = 6
_MAX_ROW_PX = 540  # Excel's row height limit is 409 pt = 545 px
_DESCRIPTION_COLUMN = 3

_CLASSIFICATION_CELLS = {"public": "B10", "commercial": "C10", "internal": "D10"}


def template_path() -> Path:
    """The template lives next to customers.json, in ``templates/``."""
    return get_settings().customers_config_path.parent / "templates" / TEMPLATE_NAME


def default_export_settings(document: Document, client_label: str) -> ExportSettings:
    """What can be deduced from the document; the user completes the rest."""
    subject = document.doc_title or Path(document.original_filename).stem
    return ExportSettings(
        client=client_label,
        reference=document.doc_reference or "",
        edition="00",
        date=date.today().isoformat(),
        subject=subject,
        doc_type="CM",  # compliance matrix
        applicable_docs=[DocumentRef(title=subject, reference=document.doc_reference or "")],
    )


def load_export_settings(document: Document, client_label: str) -> ExportSettings:
    """Saved settings, or the deduced defaults when none were saved (or they no longer validate)."""
    if document.export_settings:
        try:
            return ExportSettings.model_validate_json(document.export_settings)
        except ValidationError:
            pass
    return default_export_settings(document, client_label)


def _clean(text: str) -> str:
    """Strip the markdown emphasis and line breaks that read badly in a cell."""
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    return text.strip()[:_MAX_CELL_CHARS]


def _set_text(ws: Worksheet, ref: str | tuple[int, int], value: str) -> None:
    cell = ws[ref] if isinstance(ref, str) else ws.cell(*ref)
    cell.value = value
    if value.startswith("="):  # never let user or extracted text become a formula
        cell.data_type = "s"


def _set_if(ws: Worksheet, ref: str, value: str) -> None:
    """Overwrite a template cell only when the user provided a value."""
    if value.strip():
        _set_text(ws, ref, value.strip())


def _format_date(iso: str) -> str:
    try:
        return date.fromisoformat(iso).strftime("%d/%m/%Y")
    except ValueError:
        return date.today().strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# REQUIREMENTS sheet
# ---------------------------------------------------------------------------


def _place_picture(ws: Worksheet, row: int, picture: Picture, top_px: int) -> None:
    image = XLImage(BytesIO(picture.png))
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=_DESCRIPTION_COLUMN - 1, row=row - 1,
            colOff=pixels_to_EMU(4), rowOff=pixels_to_EMU(top_px),
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(picture.width), pixels_to_EMU(picture.height)),
    )
    ws.add_image(image)


def _fill_description(
    ws: Worksheet, row: int, markdown: str, assets_dir: Path | None, continuation_style: list
) -> int:
    """Description cell: text as text, tables and figures as pictures stacked under it.

    An Excel row holds at most 409 pt, so pictures that do not fit go to continuation rows
    right below (blank apart from the picture, so the metrics ignore them).
    Returns the number of rows used.
    """
    segments = split_segments(markdown)
    if all(seg.kind == "text" for seg in segments):
        _set_text(ws, (row, _DESCRIPTION_COLUMN), _clean(markdown))
        return 1

    texts: list[str] = []
    pictures: list[Picture] = []
    for seg in segments:
        if seg.kind == "text":
            texts.append(_clean(seg.text))
        elif seg.kind == "table":
            pictures.extend(render_table(seg.rows, styles=seg.styles))
        else:
            picture = load_picture(assets_dir / Path(seg.target).name) if assets_dir else None
            if picture:
                pictures.append(picture)
            elif seg.alt:  # figure file missing: keep its caption
                texts.append(f"[{seg.alt}]")

    text = "\n".join(t for t in texts if t)
    text_px = sum(max(1, -(-len(line) // _CHARS_PER_LINE)) for line in text.split("\n")) * _LINE_PX if text else 0
    _set_text(ws, (row, _DESCRIPTION_COLUMN), text)
    if not pictures:
        return 1

    # Greedy packing: each row takes pictures while they fit under the Excel row limit
    pages: list[list[Picture]] = [[]]
    used = text_px + _GAP_PX
    for picture in pictures:
        # a picture that does not fit under the text opens the next row, leaving the text row bare
        if used + picture.height + _GAP_PX > _MAX_ROW_PX and (pages[-1] or text_px and len(pages) == 1):
            pages.append([])
            used = _GAP_PX
        pages[-1].append(picture)
        used += picture.height + _GAP_PX

    for offset, page in enumerate(pages):
        current = row + offset
        if offset:
            for column, style in zip(_COLUMNS, continuation_style):
                ws.cell(current, column)._style = copy(style)
        if not page:  # text only: Excel fits the row height itself
            continue
        top_px = text_px + _GAP_PX if offset == 0 else _GAP_PX
        cell = ws.cell(current, _DESCRIPTION_COLUMN)
        alignment = copy(cell.alignment)
        alignment.vertical = "top"
        cell.alignment = alignment
        for picture in page:
            _place_picture(ws, current, picture, top_px)
            top_px += picture.height + _GAP_PX
        ws.row_dimensions[current].height = top_px * 0.75
    return len(pages)


def _fill_requirements(ws: Worksheet, nodes: list[Node], assets_dir: Path | None = None) -> int:
    """Write one row per node, cloning the template row styles. Returns the last row."""
    styles = {
        "title": [copy(ws.cell(_TITLE_STYLE_ROW, c)._style) for c in _COLUMNS],
        "requirement": [copy(ws.cell(_REQUIREMENT_STYLE_ROW, c)._style) for c in _COLUMNS],
    }
    for row in range(_FIRST_ROW, _EXAMPLE_LAST_ROW + 1):
        for column in _COLUMNS:
            ws.cell(row, column).value = None

    row = _FIRST_ROW
    for node in nodes:
        for column, style in zip(_COLUMNS, styles[node.type]):
            ws.cell(row, column)._style = copy(style)
        ws.row_dimensions[row].height = None  # let Excel fit the wrapped text

        used = 1
        if node.type == "title":
            # Title rows only carry the title: the text under a heading is not exported
            _set_text(ws, (row, 1), f"{node.section or ''} {node.title}".strip())
        else:
            _set_text(ws, (row, 2), node.req_id or "")
            used = _fill_description(
                ws, row, "\n".join(p for p in (node.title, node.description) if p), assets_dir,
                styles["requirement"],
            )
            if node.comment:
                _set_text(ws, (row, 6), _clean(node.comment))
        ws.cell(row, 4).value = node.applicability
        ws.cell(row, 5).value = node.conformity
        row += used
    last_row = row - 1
    return last_row


# ---------------------------------------------------------------------------
# DA-DR sheet: tables that grow by inserting rows
# ---------------------------------------------------------------------------


def _shift_names(wb, sheet: str, at_row: int, count: int) -> None:
    """Move the defined names that point below an inserted block."""
    pattern = re.compile(r"(\$[A-Z]+\$)(\d+)")
    for defined in wb.defined_names.values():
        if defined.attr_text.startswith(f"'{sheet}'!"):
            defined.attr_text = pattern.sub(
                lambda m: f"{m.group(1)}{int(m.group(2)) + count if int(m.group(2)) >= at_row else m.group(2)}",
                defined.attr_text,
            )


def _insert_rows(ws: Worksheet, at_row: int, count: int, style_row: int) -> None:
    """Insert ``count`` blank rows at ``at_row``, shifting cells, merges, heights, print area
    and defined names, then give the new rows the look of ``style_row``."""
    max_row, max_col = ws.max_row, ws.max_column
    merged = [r for r in ws.merged_cells.ranges if r.min_row >= at_row]
    template_merges = [r for r in ws.merged_cells.ranges if r.min_row == r.max_row == style_row]
    for merged_range in merged:
        ws.unmerge_cells(str(merged_range))
    if at_row <= max_row:  # nothing to shift when the table ends the sheet
        ws.move_range(f"A{at_row}:{get_column_letter(max_col)}{max_row}", rows=count)
    for merged_range in merged:
        ws.merge_cells(
            start_row=merged_range.min_row + count, end_row=merged_range.max_row + count,
            start_column=merged_range.min_col, end_column=merged_range.max_col,
        )
    for row in range(max_row, at_row - 1, -1):
        ws.row_dimensions[row + count].height = ws.row_dimensions[row].height
    for row in range(at_row, at_row + count):
        ws.row_dimensions[row].height = ws.row_dimensions[style_row].height
        for column in range(1, max_col + 1):
            ws.cell(row, column)._style = copy(ws.cell(style_row, column)._style)
        for merged_range in template_merges:
            ws.merge_cells(
                start_row=row, end_row=row,
                start_column=merged_range.min_col, end_column=merged_range.max_col,
            )
    ws.print_area = f"A1:E{int(ws.print_area.split('$')[-1]) + count}"
    _shift_names(ws.parent, ws.title, at_row, count)


def _fill_table(ws: Worksheet, first_row: int, capacity: int, rows: list[list[str]], end_name: str) -> None:
    """Write ``rows`` from ``first_row``, inserting rows when the template table is too short."""
    last_template_row = first_row + capacity - 1
    if len(rows) > capacity:
        _insert_rows(ws, last_template_row + 1, len(rows) - capacity, style_row=last_template_row)
    for offset, values in enumerate(rows):
        for column, value in enumerate(values, start=2):
            _set_text(ws, (first_row + offset, column), value)
    if rows:
        ws.parent.defined_names[end_name].attr_text = (
            f"'{ws.title}'!$B${first_row + max(len(rows), capacity) - 1}"
        )


def _fill_documents(wb, settings: ExportSettings) -> None:
    ws = wb["DA-DR"]
    # Bottom-up so the template coordinates of the upper tables stay valid
    _fill_table(
        ws, _ACRONYM_FIRST, _ACRONYM_CAPACITY,
        [[a.acronym, a.description] for a in settings.acronyms], "ACR_List_End",
    )
    _fill_table(
        ws, _RD_FIRST, _RD_CAPACITY,
        [[f"RD-{i:02d}", d.title, d.reference, d.edition] for i, d in enumerate(settings.reference_docs, 1)],
        "DR_List_End",
    )
    _fill_table(
        ws, _AD_FIRST, _AD_CAPACITY,
        [[f"AD-{i:02d}", d.title, d.reference, d.edition] for i, d in enumerate(settings.applicable_docs, 1)],
        "DA_List_End",
    )


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------


def _fill_header(wb, document: Document, settings: ExportSettings, today: str) -> None:
    header = wb["HEADER"]
    _set_if(header, "H2", settings.client)
    _set_if(header, "H3", settings.project)
    _set_if(header, "H9", settings.reference)
    header["H10"] = today
    _set_if(header, "H11", settings.edition)
    subject = settings.subject.strip() or document.doc_title or Path(document.original_filename).stem
    header["B16"] = f"Applicability and Conformity MATRIX to {subject}"
    _set_if(header, "B19", settings.subtitle)

    if settings.classification:
        for key, ref in _CLASSIFICATION_CELLS.items():
            mark = "☒" if key == settings.classification else "☐"
            cell = header[ref]
            cell.value = f"{mark} {str(cell.value).lstrip('☒☐ ')}"
            if key == settings.classification:
                font = copy(cell.font)
                font.b = True
                cell.font = font

    for ref, name, role in (
        ("D37", settings.quality_manager, "Quality Manager"),
        ("D42", settings.project_manager, "Project Manager"),
        ("D47", settings.redactor, ""),
    ):
        if name.strip():
            _set_text(header, ref, f"{name.strip()}\n{role}".strip())
    if settings.export_control:
        header["B53"] = f"EXPORT CONTROL STATUS: {settings.export_control}"


def build_compliance_matrix(
    document: Document, nodes: list[Node], settings: ExportSettings, assets_dir: Path | None = None
) -> bytes:
    """Return the filled workbook. Cells without a value keep the template placeholder."""
    wb = load_workbook(template_path())
    today = _format_date(settings.date)

    _fill_header(wb, document, settings, today)

    revisions = wb["Revisions"]
    _set_if(revisions, "A4", settings.doc_type)
    _set_if(revisions, "B11", settings.edition)
    revisions["C11"] = today
    _set_if(revisions, "D11", settings.author)
    _set_if(revisions, "F11", settings.revision_description)

    _fill_documents(wb, settings)

    requirements = wb["REQUIREMENTS"]
    last_row = _fill_requirements(requirements, nodes, assets_dir)
    print_last_row = max(last_row, int(requirements.print_area.split("$")[-1]))
    requirements.print_area = f"A1:F{print_last_row}"
    requirements.auto_filter.ref = f"A7:F{max(last_row, 7)}"

    wb.calculation.fullCalcOnLoad = True  # metrics on the Revisions sheet are formulas
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
