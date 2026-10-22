"""Tests for the Excel compliance matrix export and its settings."""

import io
from datetime import datetime
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook

from app.models.document import Document
from app.models.node import Node
from app.schemas.export import Acronym, DocumentRef, ExportSettings
from app.services.export_service import (
    build_compliance_matrix,
    default_export_settings,
    load_export_settings,
)


def _node(order: int, type_: str, **kwargs) -> Node:
    defaults = {"title": "", "description": "", "applicability": "T", "conformity": "T", "comment": ""}
    return Node(order=order, type=type_, **{**defaults, **kwargs})


@pytest.fixture
def document() -> Document:
    return Document(
        original_filename="My SOW.pdf",
        customer="airbus",
        status="completed",
        doc_title="Harness SOW",
        doc_reference="OGM-SOW-1",
        created_at=datetime(2026, 1, 1),
    )


@pytest.fixture
def nodes() -> list[Node]:
    return [
        _node(1, "title", section="4", title="Scope", description="Intro"),
        _node(2, "requirement", req_id="HAR SSOW-1", title="Inputs", description="The **Customer** delivers.",
              applicability="A/M", conformity="PC", comment="Needs a waiver"),
        _node(3, "requirement", req_id="HAR SSOW-2", title="Formula", description="=1+1",
              applicability="NA", conformity="NC"),
    ]


def _build(document, nodes, settings=None):
    settings = settings or default_export_settings(document, "Airbus")
    return load_workbook(io.BytesIO(build_compliance_matrix(document, nodes, settings)))


def test_keeps_template_sheets_and_hidden_lists(document, nodes):
    wb = _build(document, nodes)
    assert wb.sheetnames == ["HEADER", "Revisions", "DA-DR", "REQUIREMENTS", "Champs"]
    assert wb["Champs"].sheet_state == "hidden"
    assert all(len(wb[name]._images) == 1 for name in wb.sheetnames[:4])  # logos survive
    assert wb["REQUIREMENTS"].data_validations.dataValidation


def test_defaults_are_deduced_from_the_document(document, nodes):
    wb = _build(document, nodes)
    assert wb["HEADER"]["H2"].value == "Airbus"
    assert wb["HEADER"]["H9"].value == "OGM-SOW-1"
    assert wb["HEADER"]["B16"].value == "Applicability and Conformity MATRIX to Harness SOW"
    assert wb["HEADER"]["H3"].value == "PROJET"  # unknown: template placeholder kept
    assert wb["DA-DR"]["B10"].value == "AD-01"
    assert wb["DA-DR"]["C10"].value == "Harness SOW"
    assert wb["Revisions"]["A4"].value == "CM"
    assert wb["Revisions"]["D45"].value.startswith("=COUNTIF(")  # metrics stay formulas


def test_form_values_fill_the_cells(document, nodes):
    settings = ExportSettings(
        project="THOR", subtitle="Phase B2", edition="01", date="2026-03-05",
        classification="commercial", export_control="Not Listed/Not Controlled",
        quality_manager="A. Martin", project_manager="B. Durand", redactor="C. Petit",
        author="C. Petit", revision_description="Creation",
    )
    wb = _build(document, nodes, settings)
    header = wb["HEADER"]
    assert (header["H3"].value, header["H10"].value, header["H11"].value) == ("THOR", "05/03/2026", "01")
    assert header["B19"].value == "Phase B2"
    assert [header[c].value for c in ("B10", "C10", "D10")] == [
        "☐ Public", "☒ Commer. In confidence", "☐ Internal Only"
    ]
    assert header["C10"].font.b
    assert header["D37"].value == "A. Martin\nQuality Manager"
    assert header["D47"].value == "C. Petit"
    assert header["B53"].value == "EXPORT CONTROL STATUS: Not Listed/Not Controlled"
    revisions = wb["Revisions"]
    assert [revisions[c].value for c in ("B11", "C11", "D11", "F11")] == ["01", "05/03/2026", "C. Petit", "Creation"]


def test_writes_one_row_per_node_in_template_style(document, nodes):
    ws = _build(document, nodes)["REQUIREMENTS"]
    assert [ws.cell(8, c).value for c in (1, 2, 3, 4, 5, 6)] == ["4 Scope", None, None, "T", "T", None]  # titles: no description
    assert [ws.cell(9, c).value for c in (1, 2, 3, 4, 5, 6)] == [
        None, "HAR SSOW-1", "Inputs\nThe Customer delivers.", "A/M", "PC", "Needs a waiver"
    ]
    assert ws["A8"].font.b and ws["A8"].font.name == "Arial"
    assert ws["C9"].alignment.wrap_text
    assert ws.print_area.endswith("$F$105")
    assert ws.auto_filter.ref == "A7:F10"


def test_extracted_text_is_never_a_formula(document, nodes):
    cell = _build(document, nodes)["REQUIREMENTS"]["C10"]
    assert cell.data_type == "s"
    assert cell.value.endswith("=1+1")


def test_example_rows_are_cleared(document):
    ws = _build(document, [])["REQUIREMENTS"]
    assert ws["A8"].value is None and ws["B10"].value is None


def test_print_area_grows_with_the_nodes(document):
    many = [_node(i, "requirement", req_id=f"R-{i}") for i in range(150)]
    ws = _build(document, many)["REQUIREMENTS"]
    assert ws.print_area.endswith("$F$157")


def test_tables_grow_beyond_the_template_rows(document, nodes):
    settings = ExportSettings(
        applicable_docs=[DocumentRef(title=f"AD {i}", reference=f"R{i}", edition="1") for i in range(6)],
        reference_docs=[DocumentRef(title=f"RD {i}") for i in range(5)],
        acronyms=[Acronym(acronym=f"AC{i}", description=f"Desc {i}") for i in range(10)],
    )
    wb = _build(document, nodes, settings)
    ws = wb["DA-DR"]
    assert [ws.cell(r, 3).value for r in range(10, 16)] == [f"AD {i}" for i in range(6)]
    assert ws["B15"].value == "AD-06"
    # the Reference Documents section moved down by the two extra AD rows
    assert ws["A18"].value == "Documents de Référence / Reference Documents"
    assert ws["B20"].value == "DR / RD" and ws["B21"].value == "RD-01" and ws["B25"].value == "RD-05"
    assert ws["A28"].value == "Acronymes / Acronyms" or ws["A29"].value == "Acronymes / Acronyms"
    acronym_rows = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("AC")]
    assert len(acronym_rows) == 10
    assert ws.cell(acronym_rows[-1], 3).value == "Desc 9"
    # new rows inherit the template look, including the C:E merge of the acronym rows
    assert any(str(r) == f"C{acronym_rows[-1]}:E{acronym_rows[-1]}" for r in ws.merged_cells.ranges)
    assert ws.cell(15, 2).border.left.style == ws.cell(13, 2).border.left.style
    names = wb.defined_names
    assert names["DA_List_End"].attr_text == "'DA-DR'!$B$15"
    assert names["DR_List_St"].attr_text == "'DA-DR'!$B$20"
    assert names["DR_List_End"].attr_text == "'DA-DR'!$B$25"
    assert names["ACR_List_End"].attr_text == f"'DA-DR'!$B${acronym_rows[-1]}"
    assert ws.print_area.endswith(f"$E${39 + 2 + 1 + 2}")


def test_saved_settings_win_over_defaults(document):
    assert load_export_settings(document, "Airbus").client == "Airbus"
    document.export_settings = ExportSettings(project="THOR").model_dump_json()
    assert load_export_settings(document, "Airbus").project == "THOR"
    document.export_settings = "not json"
    assert load_export_settings(document, "Airbus").client == "Airbus"


async def _upload(client) -> str:
    files = {"file": ("test.pdf", io.BytesIO(b"%PDF-1.4 fake"), "application/pdf")}
    with patch("app.routers.documents.process_document", new=AsyncMock()):
        created = await client.post("/api/v1/documents", data={"customer": "airbus"}, files=files)
    return created.json()["id"]


@pytest.mark.asyncio
async def test_export_unknown_document(client):
    for path in ("export", "export-settings"):
        response = await client.get(f"/api/v1/documents/nonexistent-id/{path}")
        assert response.status_code == 404


@pytest.mark.asyncio
async def test_export_requires_processed_document(client):
    response = await client.get(f"/api/v1/documents/{await _upload(client)}/export")
    assert response.status_code == 409


@pytest.mark.asyncio
async def test_export_settings_round_trip(client):
    doc_id = await _upload(client)
    url = f"/api/v1/documents/{doc_id}/export-settings"

    defaults = (await client.get(url)).json()
    assert defaults["client"] == "Airbus" and defaults["edition"] == "00"
    assert defaults["applicable_docs"][0]["title"] == "test"

    body = {**defaults, "project": "THOR", "acronyms": [{"acronym": "SOW", "description": "Statement of Work"}]}
    assert (await client.put(url, json=body)).status_code == 200
    saved = (await client.get(url)).json()
    assert saved["project"] == "THOR" and saved["acronyms"][0]["acronym"] == "SOW"


@pytest.mark.asyncio
async def test_export_settings_reject_unknown_classification(client):
    doc_id = await _upload(client)
    response = await client.put(f"/api/v1/documents/{doc_id}/export-settings", json={"classification": "secret"})
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_update_node_comment_validation(client):
    response = await client.patch("/api/v1/nodes/some-id", json={"comment": "x" * 5001})
    assert response.status_code == 422


# ---------------------------------------------------------------------------
# Markdown tables and images in descriptions
# ---------------------------------------------------------------------------

from openpyxl.utils.units import pixels_to_EMU  # noqa: E402
from PIL import Image  # noqa: E402

from app.services.markdown_cells import render_table, split_segments  # noqa: E402

TABLE = "| Name | Value |\n| --- | --- |\n| Voltage | 28 V |\n| Current | 5 A |"


def test_split_segments_keeps_order():
    segments = split_segments(f"Before\n{TABLE}\nBetween\n![Fig 1](/api/v1/documents/d/assets/p1.png)\nAfter")
    assert [s.kind for s in segments] == ["text", "table", "text", "image", "text"]
    assert segments[1].rows == (("Name", "Value"), ("Voltage", "28 V"), ("Current", "5 A"))
    assert (segments[3].alt, segments[3].target) == ("Fig 1", "/api/v1/documents/d/assets/p1.png")


def test_a_pipe_line_without_separator_is_text():
    assert [s.kind for s in split_segments("| not | a table |")] == ["text"]


def test_render_table_fits_the_column():
    wide = (("H1", "H2"), ("x " * 200, "y"))
    for rows in ((("A", "B"), ("1", "2")), wide):
        (picture,) = render_table(rows)
        assert picture.png.startswith(b"\x89PNG")
        assert 0 < picture.width <= 481


def test_tables_become_pictures_under_the_text(document):
    node = _node(1, "requirement", req_id="R-1", title="Limits", description=f"Intro\n{TABLE}",
                 applicability="A", conformity="C")
    ws = _build(document, [node])["REQUIREMENTS"]
    assert len(ws._images) == 2  # template logo + the table
    cell = ws["C8"]
    assert cell.value == "Limits\nIntro" and "|" not in cell.value
    assert cell.alignment.vertical == "top"
    assert ws.row_dimensions[8].height > 40


def test_figures_are_embedded(document, tmp_path):
    Image.new("RGB", (800, 400), "red").save(tmp_path / "p0001_01.png")
    node = _node(1, "requirement", req_id="R-1", description="![Fig](/api/v1/documents/d/assets/p0001_01.png)")
    settings = default_export_settings(document, "Airbus")
    wb = load_workbook(io.BytesIO(build_compliance_matrix(document, [node], settings, tmp_path)))
    sizes = [(i.anchor.ext.width, i.anchor.ext.height) for i in wb["REQUIREMENTS"]._images
             if hasattr(i.anchor, "ext") and i.anchor.ext]
    assert (pixels_to_EMU(360), pixels_to_EMU(180)) in sizes  # fitted to the column, then 75%


def test_missing_figure_falls_back_to_its_caption(document, tmp_path):
    node = _node(1, "requirement", req_id="R-1", description="Text\n![Fig 2](/api/v1/documents/d/assets/gone.png)")
    ws = _build_with_assets(document, [node], tmp_path)["REQUIREMENTS"]
    assert ws["C8"].value == "Text\n[Fig 2]"
    assert len(ws._images) == 1


def test_tall_table_is_cut_in_pictures_that_each_fit_a_row():
    rows = (("A", "B"),) + tuple((f"r{i}", f"v{i}") for i in range(80))
    pictures = render_table(rows)
    assert len(pictures) > 1
    assert all(p.height <= 520 for p in pictures)


def test_tall_table_continues_on_blank_rows(document):
    big = "| A | B |\n| --- | --- |\n" + "\n".join(f"| r{i} | v{i} |" for i in range(80))
    nodes = [
        _node(1, "requirement", req_id="R-1", description=big, applicability="A", conformity="C"),
        _node(2, "requirement", req_id="R-2", description="Next", applicability="A", conformity="C"),
    ]
    ws = _build(document, nodes)["REQUIREMENTS"]
    ids = {r: ws.cell(r, 2).value for r in range(8, 16) if ws.cell(r, 2).value}
    first, second = sorted(ids)
    assert second - first > 1  # R-2 was pushed below the continuation rows
    for row in range(first + 1, second):
        assert [ws.cell(row, c).value for c in (1, 2, 3, 4, 5)] == [None] * 5
        assert ws.row_dimensions[row].height <= 409
    assert ws.row_dimensions[first].height <= 409
    assert len(ws._images) > 2
    assert ws.print_area.endswith("$F$105")


def test_tall_figure_is_scaled_to_fit_a_row(document, tmp_path):
    Image.new("RGB", (400, 3000), "green").save(tmp_path / "tall.png")
    node = _node(1, "requirement", req_id="R-1", description="![Tall](/api/v1/documents/d/assets/tall.png)")
    ws = _build_with_assets(document, [node], tmp_path)["REQUIREMENTS"]
    heights = [i.anchor.ext.height for i in ws._images if getattr(i.anchor, "ext", None)]
    assert pixels_to_EMU(390) in heights  # 520 px to fit a row, then 75%


def _build_with_assets(document, nodes, assets_dir):
    settings = default_export_settings(document, "Airbus")
    return load_workbook(io.BytesIO(build_compliance_matrix(document, nodes, settings, assets_dir)))


@pytest.mark.asyncio
async def test_asset_endpoint(client, tmp_path):
    Image.new("RGB", (10, 10), "blue").save(tmp_path / "fig.png")
    with patch("app.routers.documents.assets_dir", return_value=tmp_path):
        assert (await client.get("/api/v1/documents/d/assets/fig.png")).status_code == 200
        assert (await client.get("/api/v1/documents/d/assets/missing.png")).status_code == 404
        assert (await client.get("/api/v1/documents/d/assets/..%2Fsecret")).status_code == 404


def test_picture_that_does_not_fit_under_the_text_goes_to_the_next_row(document, tmp_path):
    Image.new("RGB", (400, 700), "red").save(tmp_path / "fig.png")
    description = ("Long intro " * 150) + "\n![Fig](/api/v1/documents/d/assets/fig.png)"
    node = _node(1, "requirement", req_id="R-1", description=description)
    ws = _build_with_assets(document, [node], tmp_path)["REQUIREMENTS"]
    assert ws.row_dimensions[8].height is None  # text row left to Excel
    assert 300 < ws.row_dimensions[9].height <= 409  # the figure sits alone on the next row
    assert ws["C9"].value is None


def test_a_row_taller_than_the_limit_is_scaled_down_and_long_words_are_broken():
    rows = (("H", "V"), ("x", "word " * 400), ("y", "Averyveryverylongunbreakablewordwithoutanyspacetobreakat" * 3))
    pictures = render_table(rows)
    assert all(p.height <= 520 and p.width <= 481 for p in pictures)


@pytest.mark.asyncio
async def test_node_description_and_comment_are_editable(client):
    from tests.conftest import _TestSession

    async with _TestSession() as session:
        doc = Document(filename="f.pdf", original_filename="f.pdf", file_path="f", customer="airbus",
                       status="completed")
        session.add(doc)
        await session.flush()
        node = Node(document_id=doc.id, order=1, type="requirement", req_id="R-1", title="T",
                    description="old", applicability="A", conformity="C")
        session.add(node)
        await session.commit()
        node_id, doc_id = node.id, doc.id

    response = await client.patch(
        f"/api/v1/nodes/{node_id}", json={"description": "| a | b |\n| --- | --- |\n| 1 | 2 |", "comment": "ok"}
    )
    assert response.status_code == 200
    assert response.json()["description"].startswith("| a | b |") and response.json()["comment"] == "ok"
    listed = (await client.get(f"/api/v1/documents/{doc_id}/nodes")).json()["data"][0]
    assert listed["description"].startswith("| a | b |")

    too_long = await client.patch(f"/api/v1/nodes/{node_id}", json={"description": "x" * 200_001})
    assert too_long.status_code == 422
